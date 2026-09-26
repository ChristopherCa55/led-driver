from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'output' / 'resume'
OUT.mkdir(parents=True, exist_ok=True)
doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.65)
sec.bottom_margin = Inches(0.65)
sec.left_margin = sec.right_margin = Inches(0.70)
sec.header_distance = sec.footer_distance = Inches(0.25)
WIDTH = 7.10

for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'List Bullet']:
    s = doc.styles[name]
    s.font.name = 'Calibri'
    s.font.color.rgb = RGBColor(0, 0, 0)
    s._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), 'Calibri')
    s._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
    for border in list(s._element.xpath('./w:pPr/w:pBdr')):
        border.getparent().remove(border)

normal = doc.styles['Normal']
normal.font.size = Pt(12)
normal.paragraph_format.line_spacing = 1.08
normal.paragraph_format.space_after = Pt(2)
normal.paragraph_format.widow_control = True

title = doc.styles['Title']
title.font.size = Pt(22)
title.font.bold = True
title.paragraph_format.space_before = Pt(0)
title.paragraph_format.space_after = Pt(3)
title.paragraph_format.line_spacing = 1

h1 = doc.styles['Heading 1']
h1.font.size = Pt(12)
h1.font.bold = True
h1.paragraph_format.space_before = Pt(14)
h1.paragraph_format.space_after = Pt(5)
h1.paragraph_format.keep_with_next = True
h1.paragraph_format.line_spacing = 1

h2 = doc.styles['Heading 2']
h2.font.size = Pt(12)
h2.font.bold = True
h2.paragraph_format.space_before = Pt(0)
h2.paragraph_format.space_after = Pt(3)
h2.paragraph_format.keep_with_next = True
h2.paragraph_format.line_spacing = 1.10

# Explicit bullet numbering gives predictable indents in Word and text extraction.
numbering = doc.part.numbering_part.element
abs_id = 20
ab = OxmlElement('w:abstractNum')
ab.set(qn('w:abstractNumId'), str(abs_id))
multilevel = OxmlElement('w:multiLevelType')
multilevel.set(qn('w:val'), 'singleLevel')
ab.append(multilevel)
lvl = OxmlElement('w:lvl')
lvl.set(qn('w:ilvl'), '0')
for tag, val in [('start','1'),('numFmt','bullet'),('lvlText','\u2022'),('lvlJc','left')]:
    el = OxmlElement('w:' + tag)
    el.set(qn('w:val'), val)
    lvl.append(el)
ppr = OxmlElement('w:pPr')
ind = OxmlElement('w:ind')
ind.set(qn('w:left'), '230')
ind.set(qn('w:hanging'), '180')
ppr.append(ind)
lvl.append(ppr)
rpr = OxmlElement('w:rPr')
fonts = OxmlElement('w:rFonts')
fonts.set(qn('w:ascii'), 'Calibri')
fonts.set(qn('w:hAnsi'), 'Calibri')
rpr.append(fonts)
lvl.append(rpr)
ab.append(lvl)
numbering.append(ab)
num = OxmlElement('w:num')
num.set(qn('w:numId'), '20')
aid = OxmlElement('w:abstractNumId')
aid.set(qn('w:val'), '20')
num.append(aid)
numbering.append(num)

def section(text):
    return doc.add_paragraph(text, 'Heading 1')

def entry(left, right, before=0):
    p = doc.add_paragraph(style='Heading 2')
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(WIDTH), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run(left)
    r = p.add_run('\t' + right)
    r.bold = False
    return p

def line(left, right=None, after=2):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.add_run(left)
    if right:
        p.paragraph_format.tab_stops.add_tab_stop(Inches(WIDTH), WD_TAB_ALIGNMENT.RIGHT)
        p.add_run('\t' + right)
    return p

def bullet(text):
    p = doc.add_paragraph(text)
    pf = p.paragraph_format
    pf.left_indent = Inches(0.16)
    pf.first_line_indent = Inches(-0.125)
    pf.space_after = Pt(5)
    pf.keep_together = True
    np = p._p.get_or_add_pPr().get_or_add_numPr()
    np.get_or_add_ilvl().val = 0
    np.get_or_add_numId().val = 20
    return p

p = doc.add_paragraph('XXXXXXXXXXX', 'Title')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = doc.add_paragraph('xxxxx@gmail.com  |  (xxx) xxx-xxxx  |  xxxx@ucsd.edu')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
for r in p.runs:
    r.font.size = Pt(10)

section('EDUCATION')
entry('University of California, San Diego', '2024 - 2028')
line('B.S. Computer Engineering, Jacobs School of Engineering', 'La Jolla, CA')
line('GPA: xxxx/4.00', after=6)
entry('California Academy of Mathematics and Science', 'Aug 2020 - Jun 2024')
line('Weighted GPA: xxx  |  Unweighted GPA: xxx', 'Carson, CA')

section('PROJECTS')
entry('SIMO LED Driver', 'Aug 2026 - Present')
bullet('Designed a three-channel single-inductor multiple-output (SIMO) boost LED driver with independent regulation and PWM dimming, progressing from LTspice simulation to KiCad PCB layout.')
bullet('Implemented peak-current control with NAND gates, Schmitt-trigger inverters, and a cross-coupled SR latch to sequence inductor charging and discharging without a dedicated controller IC.')
bullet('Modeled package and PCB parasitic inductance; tuned gate resistance to reduce simulated peak MOSFET drain-to-source voltage from 100 V to 68 V while evaluating switching losses and dead time.')

entry('Analog PWM Speed Controller', 'May 2026', before=6)
bullet('Built an analog PWM speed controller from ECE lab components, using cascaded RC filters to convert a servo PWM input to a control voltage.')
bullet('Generated a 3 kHz carrier with a CD4069 inverter oscillator and RC filter; used an LM311 comparator to produce variable-duty-cycle PWM.')

section('TECHNICAL SKILLS')
for label, text in [
    ('Design', 'LTspice, PSpice, KiCad'),
    ('Test equipment', 'Oscilloscope, multimeter'),
    ('Programming', 'Python, Java, C, MATLAB, C++'),
]:
    p = doc.add_paragraph()
    p.add_run(label + ': ').bold = True
    p.add_run(text)

section('INTERESTS')
line('RC car repair (batteries, speed controllers, motors); bicycle maintenance; photography.')

doc.core_properties.title = 'Electrical Engineering Resume'
doc.core_properties.subject = 'Education, electronics projects, and technical skills'
doc.core_properties.author = ''
doc.core_properties.last_modified_by = ''
doc.core_properties.comments = ''
doc.save(OUT / 'updated_resume.docx')
print(OUT / 'updated_resume.docx')
